"""Bounded, read-only Office text views; never render or execute document content."""

from __future__ import annotations

import io
import posixpath
import re
import zipfile
import zlib
from dataclasses import dataclass, field
from typing import Literal

from docx import Document
from docx.table import Table
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .models import APIModel

OFFICE_SUFFIXES = frozenset({".docx", ".xlsx", ".pptx"})
MAX_OFFICE_PARTS = 512
MAX_OFFICE_EXPANDED_BYTES = 16 * 1024 * 1024
MAX_OFFICE_XML_BYTES = 4 * 1024 * 1024
MAX_OFFICE_XML_NODES = 100_000
MAX_OFFICE_BLOCKS = 20_000
MAX_OFFICE_TEXT_CHARS = 2 * 1024 * 1024
MAX_OFFICE_SEARCH_BYTES = 64 * 1024 * 1024
_SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


class OfficeInvalidError(ValueError):
    pass


class OfficeLimitError(ValueError):
    pass


class OfficeBudgetError(ValueError):
    pass


@dataclass
class OfficeBudget:
    reserved_bytes: int = 0

    def reserve(self, size: int) -> None:
        if self.reserved_bytes + size > MAX_OFFICE_SEARCH_BYTES:
            raise OfficeBudgetError("Office search expansion allowance exceeded")
        self.reserved_bytes += size


class OfficeLocation(APIModel):
    kind: Literal["paragraph", "cell", "slide", "slide_notes"]
    paragraph: int | None = None
    sheet: str | None = None
    cell: str | None = None
    slide: int | None = None


@dataclass
class OfficeBlock:
    text: str
    location: OfficeLocation
    start: int


@dataclass
class OfficeText:
    notice: str
    blocks: list[OfficeBlock] = field(default_factory=list)
    length: int = 0

    def append(self, text: str, location: OfficeLocation) -> None:
        text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\v", "\n")
        if not text:
            return
        if len(self.blocks) >= MAX_OFFICE_BLOCKS or (
            self.length + len(text) + 1 > MAX_OFFICE_TEXT_CHARS
        ):
            raise OfficeLimitError("Office text extraction limit exceeded")
        self.blocks.append(OfficeBlock(text, location, self.length))
        self.length += len(text) + 1

    @property
    def text(self) -> str:
        return "\n".join(block.text for block in self.blocks)


def _inspect_package(raw: bytes, budget: OfficeBudget) -> dict[str, etree._Element]:
    """Inspect every part before library loading, including otherwise unused XML.

    No archive entry is extracted to disk. DTDs/entities and external resolution
    are disabled explicitly, including UTF-16 XML; parser defaults are not trusted.
    """
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_OFFICE_PARTS:
            raise OfficeLimitError("Office archive contains too many parts")
        names = [entry.filename for entry in entries]
        if len(set(names)) != len(names) or any(
            name.startswith("/") or "\\" in name or "\x00" in name
            or any(part in {".", ".."} for part in name.split("/")) for name in names
        ):
            raise OfficeInvalidError("Ambiguous Office archive paths")
        expanded = sum(entry.file_size for entry in entries)
        if expanded > MAX_OFFICE_EXPANDED_BYTES:
            raise OfficeLimitError("Office archive expansion exceeds the file limit")
        budget.reserve(expanded)
        trees: dict[str, etree._Element] = {}
        nodes = 0
        word_cells = 0
        for entry in entries:
            if entry.flag_bits & 1:
                raise OfficeInvalidError("Encrypted Office packages are not indexed")
            xml = entry.filename.lower().endswith((".xml", ".rels"))
            if xml and entry.file_size > MAX_OFFICE_XML_BYTES:
                raise OfficeLimitError("Office XML part exceeds the file limit")
            with archive.open(entry) as source:
                data = source.read(entry.file_size + 1)
            if len(data) != entry.file_size:
                raise OfficeInvalidError("Office part length does not match its directory")
            if not xml:
                continue
            root = etree.fromstring(data, parser=etree.XMLParser(
                resolve_entities=False, load_dtd=False, no_network=True,
                huge_tree=False, recover=False, remove_comments=True,
            ))
            if root.getroottree().docinfo.doctype:
                raise OfficeInvalidError("Office XML must not declare a DTD")
            nodes += sum(1 for _ in root.iter())
            if nodes > MAX_OFFICE_XML_NODES:
                raise OfficeLimitError("Office XML node limit exceeded")
            # python-docx expands table grid spans when exposing row.cells.
            # Bound that work before passing even a tiny crafted ZIP to it.
            word_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            for cell in root.iter(word_ns + "tc"):
                span = cell.find(f"{word_ns}tcPr/{word_ns}gridSpan")
                width = int(span.get(word_ns + "val", "1")) if span is not None else 1
                if width < 1:
                    raise OfficeInvalidError("Invalid Word table grid span")
                word_cells += width
            if word_cells > MAX_OFFICE_BLOCKS:
                raise OfficeLimitError("Word table expansion exceeds the extraction limit")
            trees[entry.filename] = root
        return trees


def _docx(raw: bytes) -> OfficeText:
    result = OfficeText(
        "Word body text and tables only; paragraph numbers follow extraction order, "
        "including table cells. Headers, footers, comments, tracked changes, text boxes, "
        "images and page layout are not represented. Verify the original before citing."
    )
    document = Document(io.BytesIO(raw))
    paragraph = 0

    def walk(container):
        nonlocal paragraph
        for item in container.iter_inner_content():
            if isinstance(item, Table):
                seen = set()
                for row in item.rows:
                    for cell in row.cells:
                        # A merged cell occurs at multiple grid positions.
                        if cell._tc not in seen:
                            seen.add(cell._tc)
                            walk(cell)
            else:
                paragraph += 1
                result.append(item.text, OfficeLocation(kind="paragraph", paragraph=paragraph))

    walk(document)
    return result


def _relationship(trees, part: str, relationship_id: str, expected_type: str) -> str:
    path = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    matches = [item for item in trees[path] if item.get("Id") == relationship_id]
    if len(matches) != 1:
        raise OfficeInvalidError("Missing or ambiguous Office relationship")
    item = matches[0]
    if (item.tag != f"{{{_PACKAGE_NS}}}Relationship" or item.get("TargetMode") == "External"
            or item.get("Type") != f"{_REL_NS}/{expected_type}"):
        raise OfficeInvalidError("Unsupported Office relationship")
    target = item.attrib["Target"]
    if "\\" in target or ":" in target or "\x00" in target:
        raise OfficeInvalidError("Invalid Office relationship target")
    resolved = posixpath.normpath(
        target.lstrip("/") if target.startswith("/")
        else posixpath.join(posixpath.dirname(part), target)
    )
    if resolved.startswith("../") or resolved not in trees:
        raise OfficeInvalidError("Office relationship leaves the package")
    return resolved


def _xlsx(trees) -> OfficeText:
    result = OfficeText(
        "Stored worksheet text and raw cell values, including hidden sheets/cells. "
        "Number/date formatting is not applied. Formulas are not executed; cached values "
        "may be stale. Charts, images, comments and external data are not indexed."
    )
    workbook = trees["xl/workbook.xml"]
    if workbook.tag != f"{{{_SHEET_NS}}}workbook":
        raise OfficeInvalidError("Unsupported spreadsheet XML namespace")
    ns = {"s": _SHEET_NS}
    shared: list[str] = []
    relationships = trees["xl/_rels/workbook.xml.rels"]
    shared_ids = [item.get("Id") for item in relationships
                  if item.get("Type") == f"{_REL_NS}/sharedStrings"]
    if len(shared_ids) > 1:
        raise OfficeInvalidError("Ambiguous shared string table")
    if shared_ids:
        path = _relationship(trees, "xl/workbook.xml", shared_ids[0], "sharedStrings")
        shared = ["".join(item.xpath("./s:t/text() | ./s:r/s:t/text()", namespaces=ns))
                  for item in trees[path].findall("s:si", ns)]
    seen_sheets = set()
    for sheet in workbook.findall("s:sheets/s:sheet", ns):
        sheet_name = sheet.attrib["name"]
        if not sheet_name or len(sheet_name) > 255 or sheet_name in seen_sheets:
            raise OfficeInvalidError("Invalid sheet name")
        seen_sheets.add(sheet_name)
        path = _relationship(trees, "xl/workbook.xml", sheet.attrib[f"{{{_REL_NS}}}id"],
                             "worksheet")
        worksheet = trees[path]
        if worksheet.tag != f"{{{_SHEET_NS}}}worksheet":
            raise OfficeInvalidError("Unsupported worksheet")
        seen = set()
        for cell in worksheet.findall("s:sheetData/s:row/s:c", ns):
            coordinate = cell.attrib["r"]
            if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", coordinate) or coordinate in seen:
                raise OfficeInvalidError("Ambiguous cell coordinates")
            column, row = re.fullmatch(r"([A-Z]+)([0-9]+)", coordinate).groups()
            if int(row) > 1_048_576 or (len(column) == 3 and column > "XFD"):
                raise OfficeInvalidError("Cell coordinates outside the worksheet")
            seen.add(coordinate)
            value = cell.findtext("s:v", default="", namespaces=ns)
            kind = cell.get("t")
            if kind == "s":
                index = int(value)
                if index < 0 or index >= len(shared):
                    raise OfficeInvalidError("Invalid shared string reference")
                value = shared[index]
            elif kind == "inlineStr":
                value = "".join(cell.xpath("./s:is/s:t/text() | ./s:is/s:r/s:t/text()",
                                          namespaces=ns))
            elif kind == "b":
                value = {"0": "FALSE", "1": "TRUE"}[value]
            formula = cell.find("s:f", ns)
            if formula is not None:
                expression = formula.text or "[expression stored elsewhere]"
                value = f"Formula: ={expression}" + (f"\nCached value: {value}" if value else "")
            result.append(value, OfficeLocation(kind="cell", sheet=sheet_name, cell=coordinate))
    return result


def _pptx(raw: bytes) -> OfficeText:
    result = OfficeText(
        "Slide text, tables and speaker notes in presentation order, including hidden slides. "
        "Images, charts, embedded objects, master/layout text and visual reading order "
        "are not represented. Verify the original before citing."
    )
    presentation = Presentation(io.BytesIO(raw))

    def texts(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from texts(shape.shapes)
            elif shape.has_text_frame:
                yield shape.text_frame.text
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if not cell.is_spanned:
                            yield cell.text

    for number, slide in enumerate(presentation.slides, start=1):
        result.append("\n".join(texts(slide.shapes)), OfficeLocation(kind="slide", slide=number))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            result.append(slide.notes_slide.notes_text_frame.text,
                          OfficeLocation(kind="slide_notes", slide=number))
    return result


def extract_office(raw: bytes, suffix: str, budget: OfficeBudget) -> OfficeText:
    try:
        trees = _inspect_package(raw, budget)
        if suffix == ".docx":
            return _docx(raw)
        if suffix == ".xlsx":
            return _xlsx(trees)
        if suffix == ".pptx":
            return _pptx(raw)
        raise OfficeInvalidError("Unsupported Office format")
    except (OfficeInvalidError, OfficeLimitError, OfficeBudgetError):
        raise
    except (ValueError, KeyError, TypeError, IndexError, AttributeError, OSError, EOFError,
            zipfile.BadZipFile, zlib.error, etree.LxmlError, RuntimeError,
            NotImplementedError) as exc:
        # Library/parser details can contain private source text; do not return them.
        raise OfficeInvalidError("Office content could not be extracted safely") from exc
