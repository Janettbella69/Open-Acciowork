from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from ophira_backend.office_export import write_docx, write_pptx, write_xlsx


def _slide_text(slide: object) -> list[str]:
    return [shape.text for shape in slide.shapes if hasattr(shape, "text")]


def test_xlsx_round_trips_typed_values_safe_text_and_wide_formula(tmp_path: Path) -> None:
    path = tmp_path / "wide.xlsx"
    headers = [f"column-{index}" for index in range(1, 29)]
    first = [None] * len(headers)
    second = [None] * len(headers)
    first[:8] = [
        "SKU-001",
        7,
        2.5,
        Decimal("3.25"),
        True,
        date(2026, 9, 2),
        "00123",
        '=HYPERLINK("https://example.test", "click")',
    ]
    first[-1] = "1.5"
    second[0] = "9" * 1_000
    second[1] = "1234567890123456"
    second[-1] = 2

    write_xlsx(path, headers, [first, second])

    workbook = load_workbook(path, data_only=False)
    sheet = workbook["Sheet1"]
    assert sheet["B2"].value == 7
    assert sheet["B2"].data_type == "n"
    assert sheet["C2"].value == 2.5
    assert sheet["C2"].data_type == "n"
    assert sheet["D2"].value == 3.25
    assert sheet["D2"].data_type == "n"
    assert sheet["E2"].value is True
    assert sheet["E2"].data_type == "b"
    assert sheet["F2"].data_type == "d"
    assert sheet["G2"].value == "00123"
    assert sheet["G2"].data_type == "s"
    assert sheet["H2"].value == '=HYPERLINK("https://example.test", "click")'
    assert sheet["H2"].data_type == "s"
    assert sheet["AB2"].value == 1.5
    assert sheet["AB2"].data_type == "n"
    assert sheet["A3"].value == "9" * 1_000
    assert sheet["A3"].data_type == "s"
    assert sheet["B3"].value == "1234567890123456"
    assert sheet["B3"].data_type == "s"
    assert sheet["AA4"].value == "Total"
    assert sheet["AB4"].value == "=SUM(AB2:AB3)"
    assert sheet["AB4"].data_type == "f"
    assert sheet.auto_filter.ref == "A1:AB3"
    assert sheet.print_title_rows == "$1:$1"
    assert sheet.print_title_cols == "$A:$A"
    assert sheet.print_area == "'Sheet1'!$A$1:$AB$4"
    workbook.close()


def test_xlsx_empty_rows_do_not_create_a_circular_total(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    write_xlsx(path, ["item", "value"], [])

    workbook = load_workbook(path, data_only=False)
    sheet = workbook.active
    assert sheet.max_row == 1
    assert sheet["A1"].value == "item"
    assert sheet["B1"].value == "value"
    assert sheet["B2"].value is None
    assert sheet.auto_filter.ref == "A1:B1"
    workbook.close()


def test_xlsx_text_last_column_does_not_create_false_total(tmp_path: Path) -> None:
    path = tmp_path / "supplier-notes.xlsx"
    write_xlsx(path, ["supplier", "status"], [["Factory A", "pending"], ["Factory B", "ready"]])

    workbook = load_workbook(path, data_only=False)
    sheet = workbook.active
    assert sheet.max_row == 3
    assert sheet["B4"].value is None
    assert sheet.print_area == "'Sheet1'!$A$1:$B$3"
    workbook.close()


def test_xlsx_cjk_headers_get_width_and_wrapped_height(tmp_path: Path) -> None:
    path = tmp_path / "headers.xlsx"
    headers = ["供应商 Supplier", "需人工核验的报价有效期和材料检测要求" * 2]
    write_xlsx(path, headers, [["工厂 A", "待核验"]])

    workbook = load_workbook(path)
    sheet = workbook.active
    assert sheet["A1"].value == headers[0]
    assert sheet["B1"].value == headers[1]
    assert sheet.column_dimensions["A"].width >= 19
    assert sheet["B1"].alignment.wrap_text
    assert sheet.row_dimensions[1].height >= 39
    workbook.close()


@pytest.mark.parametrize(
    ("headers", "rows", "message"),
    [
        ([], [], "at least one header"),
        (["only"], [[1, 2]], "more cells than headers"),
        (["x"] * 16_385, [], "16,384 columns"),
        (["value"], [["x" * 32_768]], "32,767 characters"),
    ],
)
def test_xlsx_rejects_invalid_or_lossy_inputs(
    tmp_path: Path,
    headers: list[str],
    rows: list[list[object]],
    message: str,
) -> None:
    path = tmp_path / "invalid.xlsx"
    with pytest.raises(ValueError, match=message):
        write_xlsx(path, headers, rows)
    assert not path.exists()


def test_docx_round_trips_title_and_real_paragraphs(tmp_path: Path) -> None:
    path = tmp_path / "brief.docx"
    paragraphs = ["First finding.", "", "第二段结论。", "x" * 10_000]

    write_docx(path, "Opportunity brief", paragraphs)

    document = Document(path)
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Opportunity brief",
        *paragraphs,
    ]
    assert document.paragraphs[0].style.name == "Title"
    assert document.core_properties.title == "Opportunity brief"
    title_properties = document.styles["Title"]._element.get_or_add_pPr()
    assert title_properties.find(qn("w:pBdr")) is None


def test_docx_empty_content_is_still_openable(tmp_path: Path) -> None:
    path = tmp_path / "empty.docx"
    write_docx(path, "", [])

    document = Document(path)
    assert [paragraph.text for paragraph in document.paragraphs] == [""]


def test_pptx_round_trips_text_and_paginates(tmp_path: Path) -> None:
    path = tmp_path / "pitch.pptx"
    bullets = [f"Finding {index}" for index in range(1, 13)]

    write_pptx(path, "Opportunity", bullets)

    presentation = Presentation(path)
    assert len(presentation.slides) == 3
    text_by_slide = [_slide_text(slide) for slide in presentation.slides]
    assert [sum((texts for texts in text_by_slide), []).count(item) for item in bullets] == [
        1
    ] * len(bullets)
    assert all("Opportunity" in texts for texts in text_by_slide)


def test_pptx_empty_content_is_openable_without_template_instructions(tmp_path: Path) -> None:
    path = tmp_path / "empty.pptx"
    write_pptx(path, "Empty deck", [])

    presentation = Presentation(path)
    assert len(presentation.slides) == 1
    assert _slide_text(presentation.slides[0]) == ["Empty deck"]


def test_pptx_paginates_a_long_item_without_losing_text(tmp_path: Path) -> None:
    path = tmp_path / "long.pptx"
    item = ("Detailed finding with evidence. " * 160).strip()
    write_pptx(path, "Long finding", [item])

    presentation = Presentation(path)
    assert len(presentation.slides) > 1
    fragments = [
        text
        for slide in presentation.slides
        for text in _slide_text(slide)
        if not text.isdigit() and text != "Long finding"
    ]
    assert "".join(fragments) == item
    assert all(fragment[-1].isspace() for fragment in fragments[:-1])


def test_pptx_continuations_prefer_paragraph_boundaries(tmp_path: Path) -> None:
    path = tmp_path / "paragraphs.pptx"
    item = "\n".join(
        f"Paragraph {index} " + "Review the supporting evidence. " * 6
        for index in range(14)
    )
    write_pptx(path, "Review", [item])
    presentation = Presentation(path)
    fragments = [
        shape.text.replace("\v", "\n") for slide in presentation.slides for shape in slide.shapes
        if shape.has_text_frame and shape.text != "Review" and not shape.text.isdigit()
    ]
    assert "".join(fragments) == item
    assert all(fragment.startswith("Paragraph ") for fragment in fragments)
    assert all(fragment.endswith("\n") for fragment in fragments[:-1])


def test_pptx_balances_a_lone_final_card_when_it_fits(tmp_path: Path) -> None:
    path = tmp_path / "thirteen-findings.pptx"
    bullets = [f"Finding {index}: " + "Evidence needs review. " * 8 for index in range(1, 14)]

    write_pptx(path, "Supplier review", bullets)

    presentation = Presentation(path)
    assert len(presentation.slides) == 5
    assert bullets[-2] in _slide_text(presentation.slides[-1])
    assert bullets[-1] in _slide_text(presentation.slides[-1])


@pytest.mark.parametrize("title", ["跨境供应商审核" * 36 + "终点", "W" * 255])
def test_pptx_long_title_reserves_space_and_preserves_body(tmp_path: Path, title: str) -> None:
    path = tmp_path / "long-title.pptx"
    content = "Review every source 完整保留证据 " * 90
    write_pptx(path, title, [content])

    presentation = Presentation(path)
    fragments = []
    for slide in presentation.slides:
        heading = slide.shapes[0]
        assert heading.text == title
        for shape in slide.shapes:
            if not shape.has_text_frame or shape.text == title or shape.text.isdigit():
                continue
            assert shape.top >= heading.top + heading.height + Inches(0.19)
            assert shape.top + shape.height <= Inches(6.61)
            assert shape.text_frame.paragraphs[0].alignment == PP_ALIGN.LEFT
            fragments.append(shape.text)
    assert "".join(fragments) == content


def test_pptx_excessive_title_line_breaks_fail_without_creating_a_file(tmp_path: Path) -> None:
    path = tmp_path / "invalid-title.pptx"
    with pytest.raises(ValueError, match="too many lines"):
        write_pptx(path, "line\n" * 30, ["Body"])
    assert not path.exists()


@pytest.mark.parametrize(("writer", "suffix"), [(write_docx, ".docx"), (write_pptx, ".pptx")])
def test_document_writers_reject_titles_too_long_for_core_properties(
    tmp_path: Path,
    writer: object,
    suffix: str,
) -> None:
    path = tmp_path / f"invalid{suffix}"
    with pytest.raises(ValueError, match="title cannot exceed 255 characters"):
        writer(path, "x" * 256, [])
    assert not path.exists()
