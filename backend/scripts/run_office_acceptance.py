#!/usr/bin/env python3
"""Render production Office exports and verify their downloaded bytes and content.

Uses a temporary application ledger, real LibreOffice and Poppler. No model service
or user data is involved. Page PNGs require a separate visual review; the JSON
report deliberately does not claim one. Install the backend dev dependencies and
put soffice and pdftoppm on PATH (or pass their executable paths explicitly).
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

import httpx
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT / "backend" / "src"))

from ophira_backend.config import Settings  # noqa: E402
from ophira_backend.main import create_app  # noqa: E402
from ophira_backend.models import ProjectCreate  # noqa: E402
from ophira_backend.office_export import write_docx, write_pptx, write_xlsx  # noqa: E402


def fixtures(root: Path) -> dict[str, list[str]]:
    """Expected text is based on inputs, independently of the layout algorithms."""
    expected: dict[str, list[str]] = {}
    title = "供应商审核 Supplier review"
    paragraphs = [
        "本文件使用合成验收资料。报价、交期和客户信息均不代表真实交易。",
        "The quotation requires human review before any supplier is contacted.",
        "首批 300 件，单价 12.50 USD；预计交期 30 天，运费尚待确认。",
    ]
    write_docx(root / "brief.docx", title, paragraphs)
    expected["brief.docx"] = [title, *paragraphs]

    paragraphs = [
        f"记录 {i:02d} Record {i:02d} "
        + "核对样品、材质、交期和检测证据。Review the sample and supporting evidence. " * 4
        + f"END{i:02d}"
        for i in range(1, 17)
    ]
    write_docx(root / "long-brief.docx", "完整审核记录 Complete review", paragraphs)
    expected["long-brief.docx"] = paragraphs

    findings = [
        f"Finding {i:02d}：供应商声明需要证据复核。 "
        + "Review the sample, lead time and test evidence before approval. " * 2
        + f"END{i:02d}"
        for i in range(1, 14)
    ]
    write_pptx(root / "findings.pptx", title, findings)
    expected["findings.pptx"] = findings

    long_finding = "\n".join(
        f"段落 {i:02d} " + "完整保留来源和不确定性 Evidence remains subject to review. " * 3
        + f"END{i:02d}"
        for i in range(1, 15)
    )
    write_pptx(root / "long-finding.pptx", "长文本来源记录", [long_finding])
    expected["long-finding.pptx"] = [long_finding]

    long_title = "跨境供应商审核" * 36 + "终点"
    write_pptx(root / "long-title.pptx", long_title, ["Body remains separate 正文完整保留"])
    expected["long-title.pptx"] = [long_title, "Body remains separate 正文完整保留"]

    headers = ["供应商 Supplier", "备注 Notes", "数量 Units"]
    rows = [[f"工厂 {i:02d}", "样品待核对 Sample needs review " * 3, i] for i in range(1, 31)]
    write_xlsx(root / "quantities.xlsx", headers, rows)
    expected["quantities.xlsx"] = [*headers, *(row[0] for row in rows), "465.00"]

    headers = ["记录 Record", *[f"字段 {i:02d} Field {i:02d}" for i in range(1, 12)]]
    rows = [[f"ROW{i:02d}", *[f"VALUE{i:02d}C{j:02d}" for j in range(1, 12)]]
            for i in range(1, 5)]
    write_xlsx(root / "wide.xlsx", headers, rows)
    expected["wide.xlsx"] = [*headers, *(value for row in rows for value in row)]

    rows = [["工厂 A", "待核验"], ["工厂 B", '=HYPERLINK("https://example.test","text")']]
    write_xlsx(root / "text.xlsx", ["供应商", "状态"], rows)
    expected["text.xlsx"] = [value for row in rows for value in row]
    headers = ["供应商 Supplier", "需人工核验的报价有效期和材料检测要求" * 2]
    write_xlsx(root / "long-headers.xlsx", headers, [["工厂 A", "待核验"]])
    expected["long-headers.xlsx"] = [*headers, "工厂 A", "待核验"]
    return expected


def native(command: list[str]) -> str:
    # macOS headless LO requires a stable temporary base. Never change HOME.
    env = {**os.environ, "TMPDIR": "/private/tmp"} if sys.platform == "darwin" else None
    result = subprocess.run(command, capture_output=True, text=True, timeout=90, env=env)
    if result.returncode:
        raise RuntimeError(f"Native renderer failed: {result.stdout}\n{result.stderr}")
    return result.stdout.strip()


async def verify_downloads(sources: list[Path], downloaded: Path, data: Path) -> dict[str, str]:
    app = create_app(Settings(_env_file=None, data_dir=data, anthropic_api_key=""))
    ledger = app.state.ledger
    project = ledger.create_project(ProjectCreate(name="Office acceptance", objective="Synthetic"))
    run = ledger.create_run(project.id, "Validate Office artifact versions")
    checksums = {}
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://localhost",
    ) as client:
        for source in sources:
            target = ledger.workspace_for_project(project.id) / source.name
            shutil.copyfile(source, target)
            artifact = ledger.register_artifact(project.id, source.name, run_id=run.id)
            # A later workspace change must never change a historical download.
            target.write_bytes(b"Workspace edited after registration")
            response = await client.get(artifact.download_url)
            response.raise_for_status()
            if response.content != source.read_bytes():
                raise AssertionError(f"Historical bytes changed for {source.name}")
            (downloaded / source.name).write_bytes(response.content)
            checksums[source.name] = hashlib.sha256(response.content).hexdigest()
    return checksums


def verify_recalculation(soffice: str, source: Path, root: Path) -> dict:
    root.mkdir()
    native([soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(root), str(source)])
    saved = root / source.name
    cached = load_workbook(saved, data_only=True)
    formulas = load_workbook(saved, data_only=False)
    try:
        if cached.active["C32"].value != 465 or formulas.active["C32"].value != "=SUM(C2:C31)":
            raise AssertionError("LibreOffice did not preserve and calculate the quantity total")
    finally:
        cached.close()
        formulas.close()
    return {"formula": "=SUM(C2:C31)", "cached_total": 465}


def slide_text(pdftotext: str, source: Path, pdf: Path) -> tuple[str, list[str]]:
    """Read inside each native text box, catching text that paints outside it.

    Reading per box also avoids interspersing continuation-page titles and page
    numbers into a finding that spans slides. Original fixture text is checked
    against the combined content in addition to each box's rendered text.
    """
    chunks, missing = [], []
    for index, slide in enumerate(Presentation(source).slides, start=1):
        for shape_index, shape in enumerate(slide.shapes):
            if not shape.has_text_frame or not shape.text:
                continue
            coordinates = [shape.left, shape.top, shape.width, shape.height]
            text = native([
                pdftotext, "-layout", "-f", str(index), "-l", str(index),
                *[value for key, emu in zip(["-x", "-y", "-W", "-H"], coordinates, strict=True)
                  for value in (key, str(round(emu / 12_700)))], str(pdf), "-",
            ])
            normalized = "".join(text.split())
            if "".join(shape.text.split()) not in normalized:
                missing.append(f"slide {index}: {shape.text}")
            if not shape.text.isdigit() and shape_index != 0:
                chunks.append(normalized)
    return "".join(chunks), missing


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output", required=True, type=Path, help="New acceptance bundle directory",
    )
    parser.add_argument("--soffice", default="soffice")
    parser.add_argument("--pdftoppm", default="pdftoppm")
    parser.add_argument("--pdftotext", default="pdftotext")
    parser.add_argument("--font-dir", type=Path, action="append", default=[])
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    generated, downloaded = output / "generated", output / "downloaded"
    generated.mkdir()
    downloaded.mkdir()
    if args.font_dir:
        config = ET.Element("fontconfig")
        for directory in args.font_dir:
            if not directory.is_dir():
                raise ValueError(f"Font directory does not exist: {directory}")
            ET.SubElement(config, "dir").text = str(directory.resolve())
        ET.SubElement(config, "cachedir").text = str(output / "font-cache")
        fontconfig = output / "fonts.conf"
        ET.ElementTree(config).write(fontconfig, encoding="utf-8", xml_declaration=True)
        os.environ["FONTCONFIG_FILE"] = str(fontconfig)
    report = {
        "scope": "Synthetic inputs; production Office writers and ASGI downloads; native rendering",
        "libreoffice": native([args.soffice, "--version"]),
        "font_directories": [str(directory.resolve()) for directory in args.font_dir],
        "visual_review": "pending; inspect every page PNG separately",
        "cases": [],
    }
    try:
        expected = fixtures(generated)
        with tempfile.TemporaryDirectory(prefix="ophira-office-") as data:
            checksums = asyncio.run(verify_downloads(
                sorted(generated.iterdir()), downloaded, Path(data),
            ))
        report["recalculation"] = verify_recalculation(
            args.soffice, downloaded / "quantities.xlsx", output / "recalculated",
        )
        for filename, inputs in expected.items():
            render = output / Path(filename).stem
            render.mkdir()
            native([args.soffice, "--headless", "--convert-to", "pdf", "--outdir", str(render),
                    str(downloaded / filename)])
            pdf = render / f"{Path(filename).stem}.pdf"
            pages = PdfReader(pdf).pages
            # Poppler reconstructs reading order across CJK fallback-font runs;
            # raw PDF stream order is not necessarily the visual reading order.
            text = "".join(native([args.pdftotext, "-layout", str(pdf), "-"]).split())
            # Logical reading order keeps a wrapped table header together;
            # geometric reading order can interleave a vertically centered
            # neighboring cell between its lines. Accept either text ordering.
            text += "".join(native([args.pdftotext, str(pdf), "-"]).split())
            cropped_missing = []
            if filename.endswith(".pptx"):
                body, cropped_missing = slide_text(args.pdftotext, downloaded / filename, pdf)
                text += body
            missing = [value for value in inputs if "".join(value.split()) not in text]
            missing.extend(cropped_missing)
            native([args.pdftoppm, "-scale-to", "1600", "-png", str(pdf), str(render / "page")])
            case = {"file": filename, "sha256": checksums[filename], "api_download": "passed",
                    "pages": len(pages), "rendered_content": "failed" if missing else "passed",
                    "missing_inputs": missing}
            report["cases"].append(case)
            print(f"{filename}: {len(pages)} pages; {case['rendered_content']}", flush=True)
        if any(case["missing_inputs"] for case in report["cases"]):
            raise AssertionError("Rendered content is missing; inspect report.json and page PNGs")
    finally:
        (output / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")


if __name__ == "__main__":
    main()
