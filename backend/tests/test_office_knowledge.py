from __future__ import annotations

import io
import json
import warnings
import zipfile

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from ophira_backend import office_reference
from ophira_backend.knowledge import ReferenceStatus, capture_reference, search_project_knowledge
from ophira_backend.ledger_tools import build_ledger_tools
from ophira_backend.models import ProjectCreate
from ophira_backend.office_export import write_docx, write_pptx, write_xlsx
from ophira_backend.storage import FileLedger


@pytest.fixture
def office_sources(tmp_path):
    ledger = FileLedger(tmp_path / "data")
    project = ledger.create_project(ProjectCreate(name="Office references", objective="Review"))
    return ledger, project.id


def office_bytes(tmp_path, suffix):
    path = tmp_path / f"quote{suffix}"
    if suffix == ".docx":
        write_docx(path, "Quotation", ["Acme 上海", "MOQ 500"])
    elif suffix == ".xlsx":
        write_xlsx(path, ["Supplier", "MOQ"], [["Acme 上海", 500]])
    else:
        write_pptx(path, "Quotation", ["Acme 上海", "MOQ 500"])
    return path.read_bytes()


def replace_parts(raw, updates):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED,
    ) as target:
        for entry in source.infolist():
            if entry.filename not in updates:
                target.writestr(entry.filename, source.read(entry))
        for name, data in updates.items():
            target.writestr(name, data)
    return output.getvalue()


@pytest.mark.parametrize("suffix,kind", [
    (".docx", "paragraph"), (".xlsx", "cell"), (".pptx", "slide"),
])
def test_real_office_exports_are_searchable_with_exact_sources(
    office_sources, tmp_path, suffix, kind,
):
    ledger, project_id = office_sources
    raw = office_bytes(tmp_path, suffix)
    artifact = ledger.save_upload(project_id, "quote" + suffix, raw)
    before = {path.relative_to(ledger.root): path.read_bytes()
              for path in ledger.root.rglob("*") if path.is_file()}
    result = search_project_knowledge(FileLedger(ledger.root), project_id, query="上海 Acme")
    assert result.indexed_documents == 1 and result.reference_only
    match = result.matches[0]
    assert match.artifact_id == artifact.id and match.source_location.kind == kind
    assert "Acme 上海" in match.excerpt
    assert result.documents[0].extraction_notice
    assert result.office_bytes_reserved > len(raw)
    assert capture_reference(ledger, project_id, artifact.id)[1] == raw
    after = {path.relative_to(ledger.root): path.read_bytes()
             for path in ledger.root.rglob("*") if path.is_file()}
    assert after == before


def test_docx_table_and_hyperlink_order_and_excluded_headers(office_sources, tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    path = tmp_path / "source.docx"
    write_docx(path, "Title", ["first body"])
    document = Document(path)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "table reference"
    table.cell(0, 1).text = "separate cell"
    paragraph = document.add_paragraph()
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), "local_bookmark")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "linked reference"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    document.sections[0].header.paragraphs[0].text = "header-only-canary"
    document.save(path)
    ledger, project_id = office_sources
    ledger.save_upload(project_id, path.name, path.read_bytes())
    result = search_project_knowledge(ledger, project_id, query="table")
    assert result.matches[0].source_location.paragraph == 3
    assert result.matches[0].excerpt == "table reference"
    assert "separate cell" not in result.matches[0].excerpt
    linked = search_project_knowledge(ledger, project_id, query="linked")
    assert linked.matches[0].source_location.paragraph == 5
    assert not search_project_knowledge(ledger, project_id, query="header-only-canary").matches
    assert "Headers" in result.documents[0].extraction_notice


def test_xlsx_cells_have_precise_coordinates_and_formulas_are_not_evaluated(
    office_sources, tmp_path,
):
    ledger, project_id = office_sources
    raw = office_bytes(tmp_path, ".xlsx")
    ledger.save_upload(project_id, "quote.xlsx", raw)
    result = search_project_knowledge(ledger, project_id, query="Acme")
    match = result.matches[0]
    assert match.source_location.sheet == "Sheet1" and match.source_location.cell == "A2"
    assert match.excerpt == "Acme 上海" and "500" not in match.excerpt
    # All terms still match across a document, but the excerpt has only its own locator.
    assert search_project_knowledge(ledger, project_id, query="Acme 500").matching_documents == 1
    formula = search_project_knowledge(ledger, project_id, query="SUM")
    assert formula.matches[0].source_location.cell == "B3"
    assert formula.matches[0].excerpt == "Formula: =SUM(B2:B2)"
    assert "not executed" in formula.documents[0].extraction_notice


def xlsx_parts(raw, cells, *, shared=None):
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    updates = {"xl/worksheets/sheet1.xml": (
        f'<worksheet xmlns="{ns}"><dimension ref="A1:XFD1048576"/>'
        f'<sheetData><row r="2">{cells}</row></sheetData></worksheet>'
    ).encode()}
    if shared is not None:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            rels = archive.read("xl/_rels/workbook.xml.rels").decode()
        updates["xl/_rels/workbook.xml.rels"] = rels.replace("</Relationships>", (
            '<Relationship Id="shared" Target="sharedStrings.xml" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
            'sharedStrings"/></Relationships>'
        )).encode()
        updates["xl/sharedStrings.xml"] = (
            f'<sst xmlns="{ns}"><si><r><t>{shared}</t></r><r><t> quote</t></r>'
            '<rPh sb="0" eb="1"><t>phonetic-not-source</t></rPh></si></sst>'
        ).encode()
    return replace_parts(raw, updates)


def test_sparse_xlsx_shared_strings_cached_values_and_inert_formulas(office_sources, tmp_path):
    ledger, project_id = office_sources
    raw = xlsx_parts(office_bytes(tmp_path, ".xlsx"), (
        '<c r="A2" t="s"><v>0</v></c>'
        '<c r="XFD1048576"><f>WEBSERVICE("https://private.invalid")</f><v>321</v></c>'
    ), shared="Supplier")
    ledger.save_upload(project_id, "sparse.xlsx", raw)
    result = search_project_knowledge(ledger, project_id, query="Supplier")
    assert result.matches[0].excerpt == "Supplier quote"
    assert not search_project_knowledge(ledger, project_id, query="phonetic-not-source").matches
    cached = search_project_knowledge(ledger, project_id, query="321")
    assert cached.matches[0].source_location.cell == "XFD1048576"
    assert 'Formula: =WEBSERVICE("https://private.invalid")' in cached.matches[0].excerpt
    assert "Cached value: 321" in cached.matches[0].excerpt
    assert "may be stale" in cached.documents[0].extraction_notice


def test_pptx_presentation_order_tables_group_shapes_and_notes(office_sources, tmp_path):
    path = tmp_path / "slides.pptx"
    write_pptx(path, "First slide", ["first text"])
    presentation = Presentation(path)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(3), Inches(1)).table
    table.cell(0, 0).text = "second-table-canary"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1)).text = "group-canary"
    slide.notes_slide.notes_text_frame.text = "notes-canary"
    order = presentation.slides._sldIdLst
    order.insert(0, order[-1])
    presentation.save(path)
    ledger, project_id = office_sources
    ledger.save_upload(project_id, path.name, path.read_bytes())
    result = search_project_knowledge(ledger, project_id, query="second-table-canary")
    assert result.matches[0].source_location.slide == 1
    assert "group-canary" in result.matches[0].excerpt
    assert "first text" not in result.matches[0].excerpt
    notes = search_project_knowledge(ledger, project_id, query="notes-canary")
    assert notes.matches[0].source_location.kind == "slide_notes"
    assert notes.matches[0].source_location.slide == 1
    assert notes.matches[0].excerpt == "notes-canary"


@pytest.mark.parametrize("attack", ["dtd", "utf16_dtd", "duplicate", "traversal", "corrupt"])
def test_malformed_archives_fail_closed_without_hiding_good_sources(
    office_sources, tmp_path, attack,
):
    ledger, project_id = office_sources
    raw = office_bytes(tmp_path, ".docx")
    if attack in {"dtd", "utf16_dtd"}:
        encoding = "UTF-16" if attack == "utf16_dtd" else "UTF-8"
        xml = (f'<?xml version="1.0" encoding="{encoding}"?>'
               '<!DOCTYPE root [<!ENTITY x SYSTEM "file:///private/canary">]>'
               '<root>&x;</root>')
        raw = replace_parts(raw, {"unused.xml": xml.encode(encoding)})
    elif attack == "duplicate":
        output = io.BytesIO(raw)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            with zipfile.ZipFile(output, "a") as archive:
                archive.writestr("word/document.xml", b"duplicate")
        raw = output.getvalue()
    elif attack == "traversal":
        raw = replace_parts(raw, {"../escape.xml": b"<root/>"})
    else:
        raw = b"not a ZIP"
    ledger.save_upload(project_id, "bad.docx", raw)
    good = ledger.save_upload(project_id, "good.txt", b"valid quote")
    result = search_project_knowledge(ledger, project_id, query="quote")
    assert [item.artifact_id for item in result.matches] == [good.id]
    assert next(item for item in result.documents if item.name == "bad.docx").status == (
        ReferenceStatus.INVALID_DOCUMENT
    )


@pytest.mark.parametrize("cells", [
    '<c r="A2" t="s"><v>-1</v></c>', '<c r="A2" t="s"><v>500</v></c>',
    '<c r="A2"><v>1</v></c><c r="A2"><v>2</v></c>',
    '<c r="XFE1"><v>1</v></c>', '<c r="A1048577"><v>1</v></c>',
])
def test_xlsx_ambiguous_cells_are_not_silently_indexed(office_sources, tmp_path, cells):
    ledger, project_id = office_sources
    raw = xlsx_parts(office_bytes(tmp_path, ".xlsx"), cells, shared="supplier")
    ledger.save_upload(project_id, "invalid.xlsx", raw)
    result = search_project_knowledge(ledger, project_id, query="supplier")
    assert not result.matches
    assert result.documents[0].status == ReferenceStatus.INVALID_DOCUMENT


def test_external_sheet_relationship_is_never_followed(office_sources, tmp_path):
    ledger, project_id = office_sources
    raw = office_bytes(tmp_path, ".xlsx")
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        rels = archive.read("xl/_rels/workbook.xml.rels")
    raw = replace_parts(raw, {"xl/_rels/workbook.xml.rels": rels.replace(
        b'Target="/xl/worksheets/sheet1.xml"',
        b'Target="file:///private/canary.xml" TargetMode="External"',
    )})
    ledger.save_upload(project_id, "external.xlsx", raw)
    result = search_project_knowledge(ledger, project_id, query="Acme")
    assert not result.matches and result.documents[0].status == ReferenceStatus.INVALID_DOCUMENT


@pytest.mark.parametrize("limit,value", [
    ("MAX_OFFICE_PARTS", 1), ("MAX_OFFICE_EXPANDED_BYTES", 100),
    ("MAX_OFFICE_XML_BYTES", 100), ("MAX_OFFICE_XML_NODES", 1),
    ("MAX_OFFICE_TEXT_CHARS", 5), ("MAX_OFFICE_BLOCKS", 1),
])
def test_office_extraction_limits_are_explicit(office_sources, tmp_path, monkeypatch, limit, value):
    ledger, project_id = office_sources
    raw = office_bytes(tmp_path, ".docx")
    monkeypatch.setattr(office_reference, limit, value)
    ledger.save_upload(project_id, "bounded.docx", raw)
    result = search_project_knowledge(ledger, project_id, query="Acme")
    assert not result.matches
    assert result.documents[0].status == ReferenceStatus.TOO_LARGE


def test_failed_extraction_still_consumes_shared_expansion_budget(
    office_sources, tmp_path, monkeypatch,
):
    ledger, project_id = office_sources
    raw = office_bytes(tmp_path, ".docx")
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        size = sum(item.file_size for item in archive.infolist())
    monkeypatch.setattr(office_reference, "MAX_OFFICE_SEARCH_BYTES", size + 100)
    ledger.save_upload(project_id, "older.docx", raw)
    ledger.save_upload(project_id, "newer.docx", replace_parts(raw, {"bad.xml": b"<broken>"}))
    result = search_project_knowledge(ledger, project_id, query="Acme")
    assert not result.matches
    assert [item.status for item in result.documents] == [
        ReferenceStatus.INVALID_DOCUMENT, ReferenceStatus.BUDGET_EXCEEDED,
    ]
    assert result.office_bytes_reserved == size + len(b"<broken>")


def test_word_grid_expansion_is_bounded_before_library_loading(office_sources, tmp_path):
    ledger, project_id = office_sources
    raw = replace_parts(office_bytes(tmp_path, ".docx"), {"word/document.xml": (
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b'<w:body><w:tbl><w:tr><w:tc><w:tcPr><w:gridSpan w:val="1000000000"/>'
        b'</w:tcPr><w:p/></w:tc></w:tr></w:tbl></w:body></w:document>'
    )})
    ledger.save_upload(project_id, "bomb.docx", raw)
    result = search_project_knowledge(ledger, project_id)
    assert result.documents[0].status == ReferenceStatus.TOO_LARGE


def test_changed_office_is_rejected_before_parsing(office_sources, tmp_path, monkeypatch):
    ledger, project_id = office_sources
    raw = office_bytes(tmp_path, ".docx")
    artifact = ledger.save_upload(project_id, "quote.docx", raw)
    (ledger.workspace_for_project(project_id) / artifact.path).write_bytes(raw[:-1])

    def unexpected(*args):
        pytest.fail("changed bytes must not reach Office parser")

    monkeypatch.setattr("ophira_backend.knowledge.extract_office", unexpected)
    result = search_project_knowledge(ledger, project_id, query="Acme")
    assert not result.matches and result.documents[0].status == ReferenceStatus.CHANGED


def test_office_revisions_replace_the_searchable_version(office_sources, tmp_path):
    ledger, project_id = office_sources
    old = ledger.save_upload(project_id, "quote.docx", office_bytes(tmp_path, ".docx"))
    updated = tmp_path / "updated.docx"
    write_docx(updated, "Updated", ["Replacement supplier"])
    (ledger.workspace_for_project(project_id) / old.path).write_bytes(updated.read_bytes())
    current = ledger.register_artifact(project_id, old.path)
    result = search_project_knowledge(ledger, project_id, query="Replacement")
    assert result.total_documents == 1 and result.documents[0].version == 2
    assert result.matches[0].artifact_id == current.id
    assert not search_project_knowledge(ledger, project_id, query="Acme").matches
    with pytest.raises(ValueError, match="registered version"):
        capture_reference(ledger, project_id, old.id)


@pytest.mark.asyncio
async def test_mcp_returns_office_locations_and_untrusted_reference_boundaries(
    office_sources, tmp_path,
):
    ledger, project_id = office_sources
    raw = office_bytes(tmp_path, ".xlsx")
    ledger.save_upload(project_id, "quote.xlsx", raw)
    tool = next(item for item in build_ledger_tools(ledger, project_id)
                if item.name == "search_project_knowledge")
    result = await tool.handler({"query": "Acme"})
    data = json.loads(result["content"][0]["text"])
    assert data["reference_only"]
    assert data["matches"][0]["source_location"]["cell"] == "A2"
    assert data["documents"][0]["extraction_notice"]
    assert not ledger.get_project(project_id).evidence
